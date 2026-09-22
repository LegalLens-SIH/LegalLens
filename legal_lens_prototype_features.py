from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE


OUT_PATH = r"D:\SIH\LegalLens\LEGAL_LENS_PROTOTYPE_FEATURES.pptx"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
slide = prs.slides.add_slide(prs.slide_layouts[6])

# Background
bg = slide.background.fill
bg.solid()
bg.fore_color.rgb = RGBColor(10, 18, 31)

# Top accent bar
bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.18))
bar.fill.solid()
bar.fill.fore_color.rgb = RGBColor(42, 151, 255)
bar.line.fill.background()

# Title
textbox = slide.shapes.add_textbox(Inches(0.55), Inches(0.45), Inches(8.8), Inches(0.6))
text_frame = textbox.text_frame
text_frame.word_wrap = True
paragraph = text_frame.paragraphs[0]
paragraph.text = "LEGAL LENS — PROTOTYPE FEATURES"
paragraph.font.name = "Aptos"
paragraph.font.size = Pt(24)
paragraph.font.bold = True
paragraph.font.color.rgb = RGBColor(236, 244, 255)

# Badge
badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(10.7), Inches(0.58), Inches(1.9), Inches(0.28))
badge.fill.solid()
badge.fill.fore_color.rgb = RGBColor(17, 125, 105)
badge.line.color.rgb = RGBColor(17, 125, 105)
badge_tf = badge.text_frame
badge_tf.text = "WORKING PROTOTYPE"
badge_tf.paragraphs[0].font.name = "Aptos"
badge_tf.paragraphs[0].font.size = Pt(8)
badge_tf.paragraphs[0].font.bold = True
badge_tf.paragraphs[0].font.color.rgb = RGBColor(228, 255, 250)
badge_tf.paragraphs[0].alignment = PP_ALIGN.CENTER

# Feature cards
cards = [
    {
        "title": "Multi-Photo Scanning",
        "desc": "Multiple product views\nFront / Back / Left / Right / Top\nCombined extraction",
        "color": RGBColor(59, 130, 246),
        "icon": "M",
    },
    {
        "title": "AI + OCR Analysis",
        "desc": "Image analysis\nOCR text extraction\nProduct name, MRP, quantity, origin",
        "color": RGBColor(124, 58, 237),
        "icon": "A",
    },
    {
        "title": "Compliance Engine",
        "desc": "Rule-based checks\nMissing declarations\nCompliance findings",
        "color": RGBColor(16, 185, 129),
        "icon": "C",
    },
    {
        "title": "Demo Mode",
        "desc": "No account required\nOfficial demo flow\nIsolated demo data",
        "color": RGBColor(245, 158, 11),
        "icon": "D",
    },
    {
        "title": "Dynamic Progress",
        "desc": "Image Processing → Detection → OCR → Extraction → Rule Validation → Report Generation",
        "color": RGBColor(59, 130, 246),
        "icon": "P",
    },
    {
        "title": "Compliance Score",
        "desc": "100% Fully Compliant\n90–99% Mostly Compliant\n75–89% Partially Compliant\n50–74% Action Required\n0–49% Non-Compliant\nManual Verification when score unavailable",
        "color": RGBColor(16, 185, 129),
        "icon": "S",
    },
    {
        "title": "Automated Report",
        "desc": "Structured findings\nDetected information\nMissing declarations\nApplicable rules",
        "color": RGBColor(34, 197, 94),
        "icon": "R",
    },
    {
        "title": "Scan History",
        "desc": "Stored analyses\nProduct, date, score, status\nOfficial review access",
        "color": RGBColor(14, 116, 144),
        "icon": "H",
    },
    {
        "title": "Product Repository",
        "desc": "Search and manage products\nImages, scores, scan history",
        "color": RGBColor(59, 130, 246),
        "icon": "PR",
    },
    {
        "title": "Official Dashboard",
        "desc": "Official-only access\nScanning, history, products, reports, settings, support",
        "color": RGBColor(99, 102, 241),
        "icon": "O",
    },
    {
        "title": "Unified Navigation",
        "desc": "Dashboard • Operations • History • Products • Settings • Support • Logout",
        "color": RGBColor(168, 85, 247),
        "icon": "N",
    },
    {
        "title": "Modern Responsive UI",
        "desc": "GovTech dashboard design\nResponsive cards\nConsistent colors, icons, type",
        "color": RGBColor(234, 88, 12),
        "icon": "UI",
    },
]

left = Inches(0.45)
top = Inches(1.15)
col_w = Inches(3.95)
row_h = Inches(1.28)
gap_x = Inches(0.28)
gap_y = Inches(0.18)

for index, card in enumerate(cards):
    col = index % 3
    row = index // 3
    x = left + col * (col_w + gap_x)
    y = top + row * (row_h + gap_y)

    card_shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, col_w, row_h)
    card_shape.fill.solid()
    card_shape.fill.fore_color.rgb = RGBColor(20, 31, 46)
    card_shape.line.color.rgb = card["color"]
    card_shape.line.width = Pt(1.5)

    icon_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x + Inches(0.14), y + Inches(0.12), Inches(0.42), Inches(0.42))
    icon_box.fill.solid()
    icon_box.fill.fore_color.rgb = card["color"]
    icon_box.line.color.rgb = card["color"]
    icon_tf = icon_box.text_frame
    icon_tf.clear()
    icon_p = icon_tf.paragraphs[0]
    icon_p.text = card["icon"]
    icon_p.alignment = PP_ALIGN.CENTER
    icon_p.font.name = "Aptos"
    icon_p.font.size = Pt(13)
    icon_p.font.bold = True
    icon_p.font.color.rgb = RGBColor(255, 255, 255)

    text_box = slide.shapes.add_textbox(x + Inches(0.62), y + Inches(0.08), Inches(2.9), Inches(0.95))
    text_frame = text_box.text_frame
    text_frame.word_wrap = True
    text_frame.margin_left = 0
    text_frame.margin_right = 0
    text_frame.margin_top = 0
    text_frame.margin_bottom = 0

    title_p = text_frame.paragraphs[0]
    title_p.text = card["title"]
    title_p.font.name = "Aptos"
    title_p.font.size = Pt(11)
    title_p.font.bold = True
    title_p.font.color.rgb = RGBColor(239, 246, 255)
    title_p.space_after = 1

    desc_p = text_frame.add_paragraph()
    desc_p.text = card["desc"]
    desc_p.font.name = "Aptos"
    desc_p.font.size = Pt(6.5)
    desc_p.font.color.rgb = RGBColor(191, 209, 230)
    desc_p.line_spacing = 1.1

# Bottom workflow section
workflow_y = Inches(6.5)
step_w = Inches(1.72)
start_x = Inches(0.35)
steps = [
    "MULTI-PHOTO UPLOAD",
    "AI / IMAGE ANALYSIS",
    "OCR TEXT EXTRACTION",
    "DATA STRUCTURING",
    "LEGAL METROLOGY RULE CHECK",
    "COMPLIANCE SCORE",
    "AUTOMATED REPORT",
]

for idx, step in enumerate(steps):
    x = start_x + idx * (step_w + Inches(0.06))
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, workflow_y, step_w, Inches(0.5))
    box.fill.solid()
    box.fill.fore_color.rgb = RGBColor(18, 32, 49)
    box.line.color.rgb = RGBColor(59, 130, 246)
    box.line.width = Pt(1)
    box_tf = box.text_frame
    box_tf.word_wrap = True
    box_tf.margin_left = 4
    box_tf.margin_right = 4
    box_tf.margin_top = 2
    box_tf.margin_bottom = 2
    para = box_tf.paragraphs[0]
    para.text = step
    para.alignment = PP_ALIGN.CENTER
    para.font.name = "Aptos"
    para.font.size = Pt(7.1)
    para.font.bold = True
    para.font.color.rgb = RGBColor(227, 238, 255)

    if idx < len(steps) - 1:
        arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x + step_w - Inches(0.04), workflow_y + Inches(0.16), Inches(0.24), Inches(0.18))
        arrow.fill.solid()
        arrow.fill.fore_color.rgb = RGBColor(59, 130, 246)
        arrow.line.color.rgb = RGBColor(59, 130, 246)

# Bottom divider
divider = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(7.14), Inches(12.25), Inches(0.02))
divider.fill.solid()
divider.fill.fore_color.rgb = RGBColor(63, 81, 102)
divider.line.fill.background()

prs.save(OUT_PATH)
print(f"Created: {OUT_PATH}")
